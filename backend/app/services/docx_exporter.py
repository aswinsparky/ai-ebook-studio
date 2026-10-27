import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

def clean_markdown_to_plain(text: str) -> str:
    # Basic markdown stripper for docx inline formatting
    clean = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    clean = re.sub(r"\*(.*?)\*", r"\1", clean)
    clean = re.sub(r"`(.*?)`", r"\1", clean)
    clean = re.sub(r"^>\s*", "", clean)
    return clean

def generate_docx(book, chapters, settings, output_path: str):
    doc = Document()

    # Set Margins
    margin_val = settings.margin_inches if settings else 0.75
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(margin_val)
        section.bottom_margin = Inches(margin_val)
        section.left_margin = Inches(margin_val)
        section.right_margin = Inches(margin_val)

    # Title Page
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(120)
    title_run = title_p.add_run(book.title)
    title_run.font.name = settings.font_family if settings else "Georgia"
    title_run.font.size = Pt(32)
    title_run.bold = True
    title_run.font.color.rgb = RGBColor(0x11, 0x18, 0x27)

    if book.subtitle:
        sub_p = doc.add_paragraph()
        sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub_p.paragraph_format.space_before = Pt(12)
        sub_run = sub_p.add_run(book.subtitle)
        sub_run.font.name = settings.font_family if settings else "Georgia"
        sub_run.font.size = Pt(16)
        sub_run.italic = True
        sub_run.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

    author_p = doc.add_paragraph()
    author_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author_p.paragraph_format.space_before = Pt(40)
    author_run = author_p.add_run(f"By {book.author or 'Anonymous'}")
    author_run.font.name = settings.font_family if settings else "Georgia"
    author_run.font.size = Pt(14)
    author_run.bold = True

    doc.add_page_break()

    # Front Matter (Intro, Preface, Acknowledgements)
    if settings and settings.include_front_matter and settings.front_matter_json:
        import json
        try:
            fm = json.loads(settings.front_matter_json)
            for key, content in fm.items():
                if content:
                    sec_heading = doc.add_heading(key.replace("_", " ").title(), level=1)
                    sec_heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for line in content.split("\n\n"):
                        clean_line = clean_markdown_to_plain(line.strip())
                        if clean_line:
                            p = doc.add_paragraph(clean_line)
                            p.paragraph_format.line_spacing = settings.line_spacing or 1.4
                    doc.add_page_break()
        except Exception:
            pass

    # Table of Contents placeholder
    if settings and settings.include_toc:
        toc_heading = doc.add_heading("Table of Contents", level=1)
        for ch in chapters:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(f"Chapter {ch.order_index}: {ch.title}")
            run.font.name = settings.font_family if settings else "Georgia"
            run.font.size = Pt(11)
        doc.add_page_break()

    # Chapters
    for ch in chapters:
        ch_heading = doc.add_heading(f"Chapter {ch.order_index}: {ch.title}", level=1)
        ch_heading.paragraph_format.space_before = Pt(24)
        ch_heading.paragraph_format.space_after = Pt(18)

        content = ch.content or "No content recorded."
        paragraphs = content.split("\n\n")
        for para in paragraphs:
            trimmed = para.strip()
            if not trimmed:
                continue

            if trimmed.startswith("### "):
                subh = doc.add_heading(clean_markdown_to_plain(trimmed[4:]), level=3)
                subh.paragraph_format.space_before = Pt(12)
            elif trimmed.startswith("## "):
                subh = doc.add_heading(clean_markdown_to_plain(trimmed[3:]), level=2)
                subh.paragraph_format.space_before = Pt(16)
            elif trimmed.startswith("# "):
                continue  # Skip duplicate chapter H1
            elif trimmed.startswith("- ") or trimmed.startswith("* "):
                for item in trimmed.split("\n"):
                    item_text = re.sub(r"^[\-\*]\s*", "", item.strip())
                    doc.add_paragraph(clean_markdown_to_plain(item_text), style='List Bullet')
            else:
                p = doc.add_paragraph()
                p.paragraph_format.line_spacing = settings.line_spacing if settings else 1.4
                p.paragraph_format.space_after = Pt(8)
                p.add_run(clean_markdown_to_plain(trimmed))

        doc.add_page_break()

    # Back Matter (Conclusion, Author Bio, Glossary)
    if settings and settings.include_back_matter and settings.back_matter_json:
        import json
        try:
            bm = json.loads(settings.back_matter_json)
            for key, content in bm.items():
                if content:
                    sec_heading = doc.add_heading(key.replace("_", " ").title(), level=1)
                    sec_heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for line in content.split("\n\n"):
                        clean_line = clean_markdown_to_plain(line.strip())
                        if clean_line:
                            p = doc.add_paragraph(clean_line)
                            p.paragraph_format.line_spacing = settings.line_spacing or 1.4
                    doc.add_page_break()
        except Exception:
            pass

    doc.save(output_path)
    return output_path

